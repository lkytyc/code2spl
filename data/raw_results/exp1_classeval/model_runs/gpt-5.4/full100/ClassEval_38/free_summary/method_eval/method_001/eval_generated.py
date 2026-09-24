from openpyxl import Workbook, load_workbook


class ExcelProcessor:
    def __init__(self):
        pass

    def read_excel(self, file_name):
        try:
            workbook = load_workbook(file_name)
            sheet = workbook.active
            data = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
            return data
        except Exception:
            return None

    def write_excel(self, data, file_name):
        workbook = None
        try:
            workbook = Workbook()
            sheet = workbook.active
            for row in data:
                sheet.append(list(row))
            workbook.save(file_name)
            workbook.close()
            return 1
        except Exception:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            return 0

    def process_excel_data(self, N, save_file_name):
        data = self.read_excel(save_file_name)
        if data is None or not data or N < 0 or N >= len(data[0]):
            return 0

        processed_data = []
        for row in data:
            row_list = list(row)
            value = row_list[N]
            value_str = str(value)
            if not value_str.isdigit():
                row_list.append(value_str.upper())
            else:
                row_list.append(value)
            processed_data.append(tuple(row_list))

        new_file_name = save_file_name.split('.')[0] + '_process.xlsx'
        result = self.write_excel(processed_data, new_file_name)
        return result, new_file_name

import unittest
import os

class ExcelProcessorTestWriteExcel(unittest.TestCase):
    def test_write_excel_1(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia'),
            ('Julia', 28, 'Germany')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_2(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age'),
            ('John', 25),
            ('Alice', 30),
            ('Bob', 35),
            ('Julia', 28)
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_3(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_4(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_5(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_6(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA')
        ]
        save_file_name = ''
        success = processor.write_excel(new_data, save_file_name)
        self.assertEqual(success, 0)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
