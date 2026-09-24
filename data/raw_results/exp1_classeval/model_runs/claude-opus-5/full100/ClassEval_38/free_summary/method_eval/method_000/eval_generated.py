import openpyxl


class ExcelProcessor:

    def read_excel(self, file_name):
        try:
            workbook = openpyxl.load_workbook(file_name)
            sheet = workbook.active
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append(tuple(row))
            return data
        except:
            return None

    def write_excel(self, data, file_name):
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for row in data:
                sheet.append(list(row))
            workbook.save(file_name)
            return 1
        except:
            return 0

    def process_excel_data(self, N, save_file_name):
        data = self.read_excel(save_file_name)

        if data is None:
            return 0

        if len(data) == 0:
            return 0

        if N < 0 or N >= len(data[0]):
            return 0

        processed_rows = []
        for row in data:
            new_row = list(row)
            cell_value = row[N]
            cell_str = str(cell_value)
            if not cell_str.isdigit():
                new_row.append(cell_str.upper())
            else:
                new_row.append(cell_value)
            processed_rows.append(tuple(new_row))

        parts = save_file_name.split('.')
        new_file_name = parts[0] + '_process.' + parts[1]

        success = self.write_excel(processed_rows, new_file_name)
        return (success, new_file_name)

import unittest
import os

class ExcelProcessorTestReadExcel(unittest.TestCase):
    def test_read_excel_1(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_2(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age'],
                ['John', 25],
                ['Alice', 30],
                ['Bob', 35]]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age'),
            ('John', 25),
            ('Alice', 30),
            ('Bob', 35)
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_3(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name'],
                ['John'],
                ['Alice'],
                ['Bob']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name',),
            ('John',),
            ('Alice',),
            ('Bob',)
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_4(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Country'],
                ['John', 'USA'],
                ['Alice', 'Canada'],
                ['Bob', 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Country'),
            ('John', 'USA'),
            ('Alice', 'Canada'),
            ('Bob', 'Australia')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_5(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Country'],
                ['John', 'USA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Country'),
            ('John', 'USA')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_6(self):
        self.test_file_name = ''
        processor = ExcelProcessor()
        res = processor.read_excel(self.test_file_name)
        self.assertEqual(res, None)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
