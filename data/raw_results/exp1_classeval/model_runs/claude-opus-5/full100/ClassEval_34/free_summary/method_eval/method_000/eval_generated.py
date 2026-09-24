from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocFileHandler:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def _get_alignment_value(self, alignment: str):
        mapping = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
        }
        return mapping.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def read_text(self) -> str:
        doc = Document(self.file_path)
        paragraphs = [para.text for para in doc.paragraphs]
        return '\n'.join(paragraphs)

    def write_text(self, text: str, font_size: int = 12, alignment: str = 'left') -> bool:
        try:
            doc = Document()
            paragraph = doc.add_paragraph(text)
            paragraph.alignment = self._get_alignment_value(alignment)
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
            doc.save(self.file_path)
            return True
        except:
            return False

    def add_heading(self, text: str, level: int = 1) -> bool:
        try:
            doc = Document(self.file_path)
            doc.add_heading(text, level=level)
            doc.save(self.file_path)
            return True
        except:
            return False

    def add_table(self, data: list) -> bool:
        try:
            doc = Document(self.file_path)
            rows = len(data)
            cols = len(data[0]) if rows > 0 else 0
            table = doc.add_table(rows=rows, cols=cols)
            for i, row in enumerate(data):
                for j, value in enumerate(row):
                    table.cell(i, j).text = str(value)
            doc.save(self.file_path)
            return True
        except:
            return False

import unittest
import os

class DocFileHandlerTestReadText(unittest.TestCase):
    def test_read_text_1(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "Initial content"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_2(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("111")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "111"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_3(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("aaa")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "aaa"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_4(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("aaa\nbbb")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "aaa\nbbb"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_5(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = ""
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
