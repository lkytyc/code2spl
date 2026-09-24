from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocFileHandler:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def _get_alignment_value(self, alignment: str):
        mapping = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
        }
        return mapping.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def read_text(self) -> str:
        doc = Document(self.file_path)
        paragraphs = [para.text for para in doc.paragraphs]
        return '\n'.join(paragraphs)

    def write_text(self, text: str, font_size: int = 12, alignment: str = 'left') -> bool:
        try:
            doc = Document()
            paragraph = doc.add_paragraph(text)
            paragraph.alignment = self._get_alignment_value(alignment)
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
            doc.save(self.file_path)
            return True
        except:
            return False

    def add_heading(self, text: str, level: int = 1) -> bool:
        try:
            doc = Document(self.file_path)
            doc.add_heading(text, level=level)
            doc.save(self.file_path)
            return True
        except:
            return False

    def add_table(self, data: list) -> bool:
        try:
            doc = Document(self.file_path)
            rows = len(data)
            cols = len(data[0]) if rows > 0 else 0
            table = doc.add_table(rows=rows, cols=cols)
            for i, row in enumerate(data):
                for j, value in enumerate(row):
                    table.cell(i, j).text = str(value)
            doc.save(self.file_path)
            return True
        except:
            return False

import unittest
import os

class DocFileHandlerTestAddTable(unittest.TestCase):
    def setUp(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

    def tearDown(self):
        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_add_table_1(self):
        data = [['Name', 'Age']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 1)
        self.assertEqual(len(table.columns), 2)

    def test_add_table_2(self):
        data = [['Name', 'Age'], ['John', '25']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'John')

    def test_add_table_3(self):
        data = [['Name', 'Age'], ['John', '25'], ['Emma', '30']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'John')
        self.assertEqual(table.cell(2, 1).text, '30')

    def test_add_table_4(self):
        data = [['Name', 'Age'], ['aaa', '25'], ['Emma', '30']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'aaa')
        self.assertEqual(table.cell(2, 1).text, '30')

    def test_add_table_5(self):
        data = [['Name', 'Age'], ['John', '25'], ['Emma', '90']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'John')
        self.assertEqual(table.cell(2, 1).text, '90')

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
