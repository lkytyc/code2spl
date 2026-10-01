from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class DocFileHandler:
    def __init__(self, file_path: str):
        self.file_path = file_path

    def _get_alignment_value(self, alignment: str) -> WD_PARAGRAPH_ALIGNMENT:
        alignment_options = {
            'left': WD_PARAGRAPH_ALIGNMENT.LEFT,
            'center': WD_PARAGRAPH_ALIGNMENT.CENTER,
            'right': WD_PARAGRAPH_ALIGNMENT.RIGHT,
        }
        normalized_alignment = alignment.lower()
        alignment_value = alignment_options.get(normalized_alignment, WD_PARAGRAPH_ALIGNMENT.LEFT)
        return alignment_value

    def add_heading(self, heading: str, level: int = 1) -> bool:
        try:
            doc = Document(self.file_path)
            doc.add_heading(heading, level=level)
            doc.save(self.file_path)
            return True
        except:
            return False

    def add_table(self, data: list) -> bool:
        try:
            doc = Document(self.file_path)
            table = doc.add_table(rows=len(data), cols=len(data[0]))
            for i in range(len(data)):
                for j in range(len(data[i])):
                    cell_value = data[i][j]
                    table.cell(i, j).text = str(cell_value)
            doc.save(self.file_path)
            return True
        except:
            return False

    def read_text(self) -> str:
        doc = Document(self.file_path)
        text = []
        for paragraph in doc.paragraphs:
            text.append(paragraph.text)
        return '\n'.join(text)

    def write_text(self, content: str, font_size: float = 12, alignment: str = 'left') -> bool:
        try:
            doc = Document()
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(content)
            font = run.font
            font.size = Pt(font_size)
            alignment_value = self._get_alignment_value(alignment)
            paragraph.alignment = alignment_value
            doc.save(self.file_path)
            return True
        except:
            return False

import unittest
import os

class DocFileHandlerTestAddTable(unittest.TestCase):
    def setUp(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

    def tearDown(self):
        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_add_table_1(self):
        data = [['Name', 'Age']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 1)
        self.assertEqual(len(table.columns), 2)

    def test_add_table_2(self):
        data = [['Name', 'Age'], ['John', '25']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'John')

    def test_add_table_3(self):
        data = [['Name', 'Age'], ['John', '25'], ['Emma', '30']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'John')
        self.assertEqual(table.cell(2, 1).text, '30')

    def test_add_table_4(self):
        data = [['Name', 'Age'], ['aaa', '25'], ['Emma', '30']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'aaa')
        self.assertEqual(table.cell(2, 1).text, '30')

    def test_add_table_5(self):
        data = [['Name', 'Age'], ['John', '25'], ['Emma', '90']]
        self.handler.add_table(data)
        doc = Document(self.file_path)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        self.assertEqual(table.cell(1, 0).text, 'John')
        self.assertEqual(table.cell(2, 1).text, '90')

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
