class ExcelProcessor:
    def __init__(self):
        try:
            import openpyxl
            self.openpyxl = openpyxl
        except ImportError:
            self.openpyxl = None

    def read_excel(self, file_name):
        if self.openpyxl is None:
            return None
        try:
            wb = self.openpyxl.load_workbook(file_name)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(tuple(row))
            wb.close()
            return rows
        except Exception:
            return None

    def write_excel(self, data, file_name):
        if self.openpyxl is None:
            return 0
        try:
            wb = self.openpyxl.Workbook()
            ws = wb.active
            for row in data:
                ws.append(list(row))
            wb.save(file_name)
            wb.close()
            return 1
        except Exception:
            return 0

    def process_excel_data(self, N, save_file_name):
        rows = self.read_excel(save_file_name)
        if rows is None:
            return 0
        if not rows:
            return 0
        if N < 0 or N >= len(rows[0]):
            return 0
        processed = []
        for row in rows:
            row = list(row)
            value = row[N]
            if value is not None and str(value).isdigit():
                row.append(value)
            else:
                row.append(str(value).upper())
            processed.append(row)
        new_file_name = save_file_name.replace('.xlsx', '_process.xlsx')
        if new_file_name == save_file_name:
            new_file_name = save_file_name + '_process'
        result = self.write_excel(processed, new_file_name)
        return (result, new_file_name)

import unittest
import os

class ExcelProcessorTestProcessExcelData(unittest.TestCase):
    def test_process_excel_data_1(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'Canada', 30),
            ('Bob', 35, 'Australia', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_2(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 0
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'NAME'),
            ('John', 25, 'USA', 'JOHN'),
            ('Alice', 30, 'Canada', 'ALICE'),
            ('Bob', 35, 'Australia', 'BOB')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_3(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 2
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'COUNTRY'),
            ('John', 25, 'USA', 'USA'),
            ('Alice', 30, 'Canada', 'CANADA'),
            ('Bob', 35, 'Australia', 'AUSTRALIA')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_4(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 2
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'COUNTRY', 'COUNTRY'),
            ('John', 25, 'USA', 'USA'),
            ('Alice', 30, 'CANADA', 'CANADA'),
            ('Bob', 35, 'AUSTRALIA', 'AUSTRALIA')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_5(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'AGE', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'AGE', 'COUNTRY', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'CANADA', 30),
            ('Bob', 35, 'AUSTRALIA', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_6(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'AGE', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        res = processor.process_excel_data(100, self.test_file_name)
        self.assertEqual(res, 0)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
