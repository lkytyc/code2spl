from openpyxl import Workbook, load_workbook

class ExcelProcessor:
    def read_excel(self, file_name):
        try:
            workbook = load_workbook(file_name)
            sheet = workbook.active
            return [row for row in sheet.iter_rows(values_only=True)]
        except Exception:
            return None

    def write_excel(self, data, file_name):
        try:
            workbook = Workbook()
            sheet = workbook.active
            for row in data:
                sheet.append(row)
            workbook.save(file_name)
            return 1
        except Exception:
            return 0

    def process_excel_data(self, N, save_file_name):
        data = self.read_excel(save_file_name)
        if data is None or not data or N < 0 or N >= len(data[0]):
            return 0

        new_data = []
        for row in data:
            value = row[N]
            if not str(value).isdigit():
                new_row = list(row) + [str(value).upper()]
            else:
                new_row = list(row) + [value]
            new_data.append(new_row)

        new_file_name = save_file_name.replace(".xlsx", "_process.xlsx")
        success = self.write_excel(new_data, new_file_name)
        return (success, new_file_name)

import unittest
import os


class ExcelProcessorTestReadExcel(unittest.TestCase):
    def test_read_excel_1(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_2(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age'],
                ['John', 25],
                ['Alice', 30],
                ['Bob', 35]]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age'),
            ('John', 25),
            ('Alice', 30),
            ('Bob', 35)
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_3(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name'],
                ['John'],
                ['Alice'],
                ['Bob']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name',),
            ('John',),
            ('Alice',),
            ('Bob',)
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_4(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Country'],
                ['John', 'USA'],
                ['Alice', 'Canada'],
                ['Bob', 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Country'),
            ('John', 'USA'),
            ('Alice', 'Canada'),
            ('Bob', 'Australia')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_5(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Country'],
                ['John', 'USA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Country'),
            ('John', 'USA')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_6(self):
        self.test_file_name = ''
        processor = ExcelProcessor()
        res = processor.read_excel(self.test_file_name)
        self.assertEqual(res, None)


class ExcelProcessorTestWriteExcel(unittest.TestCase):
    def test_write_excel_1(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia'),
            ('Julia', 28, 'Germany')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_2(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age'),
            ('John', 25),
            ('Alice', 30),
            ('Bob', 35),
            ('Julia', 28)
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_3(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_4(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_5(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

    def test_write_excel_6(self):
        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA')
        ]
        save_file_name = ''
        success = processor.write_excel(new_data, save_file_name)
        self.assertEqual(success, 0)


class ExcelProcessorTestProcessExcelData(unittest.TestCase):
    def test_process_excel_data_1(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'Canada', 30),
            ('Bob', 35, 'Australia', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_2(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 0
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'NAME'),
            ('John', 25, 'USA', 'JOHN'),
            ('Alice', 30, 'Canada', 'ALICE'),
            ('Bob', 35, 'Australia', 'BOB')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_3(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 2
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'COUNTRY'),
            ('John', 25, 'USA', 'USA'),
            ('Alice', 30, 'Canada', 'CANADA'),
            ('Bob', 35, 'Australia', 'AUSTRALIA')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_4(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 2
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'COUNTRY', 'COUNTRY'),
            ('John', 25, 'USA', 'USA'),
            ('Alice', 30, 'CANADA', 'CANADA'),
            ('Bob', 35, 'AUSTRALIA', 'AUSTRALIA')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_5(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'AGE', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'AGE', 'COUNTRY', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'CANADA', 30),
            ('Bob', 35, 'AUSTRALIA', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_6(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'AGE', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        res = processor.process_excel_data(100, self.test_file_name)
        self.assertEqual(res, 0)


class ExcelProcessorTest(unittest.TestCase):
    def test_ExcelProcessor(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia')
        ]
        self.assertEqual(data, expected_data)

        processor = ExcelProcessor()
        new_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia'),
            ('Julia', 28, 'Germany')
        ]
        save_file_name = 'test_output.xlsx'
        success = processor.write_excel(new_data, save_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.exists(save_file_name))
        saved_data = processor.read_excel(save_file_name)
        self.assertEqual(saved_data, new_data)
        os.remove(save_file_name)

        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'Canada', 30),
            ('Bob', 35, 'Australia', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
