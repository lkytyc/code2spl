import openpyxl


class ExcelProcessor:
    def __init__(self):
        pass

    def read_excel(self, file_name):
        data = []
        try:
            workbook = openpyxl.load_workbook(file_name)
            sheet = workbook.active
            for row in sheet.iter_rows(values_only=True):
                data.append(row)
            workbook.close()
            return data
        except:
            return None

    def write_excel(self, data, file_name):
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for row in data:
                sheet.append(row)
            workbook.save(file_name)
            workbook.close()
            return 1
        except:
            return 0

    def process_excel_data(self, N, save_file_name):
        data = self.read_excel(save_file_name)
        if data is None or not data or N >= len(data[0]):
            return 0
        new_data = []
        for row in data:
            new_row = list(row[:])
            if not str(row[N]).isdigit():
                new_row.append(str(row[N]).upper())
            else:
                new_row.append(row[N])
            new_data.append(new_row)
        new_file_name = save_file_name.split('.')[0] + '_process.xlsx'
        success = self.write_excel(new_data, new_file_name)
        return success, new_file_name

import unittest
import os

class ExcelProcessorTestProcessExcelData(unittest.TestCase):
    def test_process_excel_data_1(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'Canada', 30),
            ('Bob', 35, 'Australia', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_2(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 0
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'NAME'),
            ('John', 25, 'USA', 'JOHN'),
            ('Alice', 30, 'Canada', 'ALICE'),
            ('Bob', 35, 'Australia', 'BOB')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_3(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 2
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'Country', 'COUNTRY'),
            ('John', 25, 'USA', 'USA'),
            ('Alice', 30, 'Canada', 'CANADA'),
            ('Bob', 35, 'Australia', 'AUSTRALIA')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_4(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 2
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'Age', 'COUNTRY', 'COUNTRY'),
            ('John', 25, 'USA', 'USA'),
            ('Alice', 30, 'CANADA', 'CANADA'),
            ('Bob', 35, 'AUSTRALIA', 'AUSTRALIA')
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_5(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'AGE', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        N = 1
        success, output_file = processor.process_excel_data(N, self.test_file_name)
        self.assertTrue(success)
        self.assertTrue(os.path.isfile(output_file))
        processed_data = processor.read_excel(output_file)
        expected_processed_data = [
            ('Name', 'AGE', 'COUNTRY', 'AGE'),
            ('John', 25, 'USA', 25),
            ('Alice', 30, 'CANADA', 30),
            ('Bob', 35, 'AUSTRALIA', 35)
        ]
        self.assertEqual(processed_data, expected_processed_data)
        os.remove(output_file)

    def test_process_excel_data_6(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'AGE', 'COUNTRY'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'CANADA'],
                ['Bob', 35, 'AUSTRALIA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        res = processor.process_excel_data(100, self.test_file_name)
        self.assertEqual(res, 0)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
