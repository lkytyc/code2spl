from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


class DocFileHandler:
    def __init__(self, file_path):
        self.file_path = file_path

    def read_text(self):
        doc = Document(self.file_path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    def write_text(self, content, font_size=None, alignment="left"):
        try:
            doc = Document()
            paragraph = doc.add_paragraph(content)
            paragraph.alignment = self._map_alignment(alignment)
            if font_size is not None:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def add_heading(self, text, level=1):
        try:
            doc = Document(self.file_path)
            doc.add_heading(text, level=level)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def add_table(self, data):
        try:
            doc = Document(self.file_path)
            if data:
                rows = len(data)
                cols = max(len(row) for row in data)
            else:
                rows = 0
                cols = 0
            table = doc.add_table(rows=rows, cols=cols)
            for r, row in enumerate(data):
                for c, value in enumerate(row):
                    table.cell(r, c).text = str(value)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def _map_alignment(self, alignment):
        if isinstance(alignment, str):
            mapping = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
            }
            return mapping.get(alignment.lower(), WD_ALIGN_PARAGRAPH.LEFT)
        return WD_ALIGN_PARAGRAPH.LEFT

import unittest
import os

class DocFileHandlerTestWriteText(unittest.TestCase):
    def setUp(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

    def tearDown(self):
        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_write_text_1(self):
        new_content = "New content 1"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_2(self):
        new_content = "New content 2"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_3(self):
        new_content = "New content 3"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_4(self):
        new_content = "New content 4"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_5(self):
        new_content = ""
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
