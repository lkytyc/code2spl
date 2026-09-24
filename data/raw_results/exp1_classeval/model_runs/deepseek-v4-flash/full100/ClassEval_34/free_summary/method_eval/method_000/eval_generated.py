from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


class DocFileHandler:
    def __init__(self, file_path):
        self.file_path = file_path

    def read_text(self):
        doc = Document(self.file_path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    def write_text(self, content, font_size=None, alignment="left"):
        try:
            doc = Document()
            paragraph = doc.add_paragraph(content)
            paragraph.alignment = self._map_alignment(alignment)
            if font_size is not None:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def add_heading(self, text, level=1):
        try:
            doc = Document(self.file_path)
            doc.add_heading(text, level=level)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def add_table(self, data):
        try:
            doc = Document(self.file_path)
            if data:
                rows = len(data)
                cols = max(len(row) for row in data)
            else:
                rows = 0
                cols = 0
            table = doc.add_table(rows=rows, cols=cols)
            for r, row in enumerate(data):
                for c, value in enumerate(row):
                    table.cell(r, c).text = str(value)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def _map_alignment(self, alignment):
        if isinstance(alignment, str):
            mapping = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
            }
            return mapping.get(alignment.lower(), WD_ALIGN_PARAGRAPH.LEFT)
        return WD_ALIGN_PARAGRAPH.LEFT

import unittest
import os

class DocFileHandlerTestReadText(unittest.TestCase):
    def test_read_text_1(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "Initial content"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_2(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("111")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "111"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_3(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("aaa")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "aaa"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_4(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("aaa\nbbb")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = "aaa\nbbb"
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_read_text_5(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("")
        doc.save(self.file_path)

        text_content = self.handler.read_text()
        expected_content = ""
        self.assertEqual(text_content, expected_content)

        if os.path.exists(self.file_path):
            os.remove(self.file_path)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
