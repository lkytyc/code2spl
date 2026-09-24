class ExcelProcessor:
    def read_excel(self, file_name):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(file_name)
            sheet = workbook.active
            return list(sheet.iter_rows(values_only=True))
        except Exception:
            return None

    def write_excel(self, data, file_name):
        try:
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            for row in data:
                sheet.append(list(row))
            workbook.save(file_name)
            return 1
        except Exception:
            return 0

    def process_excel_data(self, N, save_file_name):
        data = self.read_excel(save_file_name)
        if not data:
            return 0

        if not isinstance(N, int) or N < 0 or N >= len(data[0]):
            return 0

        processed_data = []
        for row in data:
            row_list = list(row)
            if N >= len(row_list):
                return 0
            value = row_list[N]
            if isinstance(value, str):
                is_digits = value.isdigit()
            else:
                is_digits = str(value).isdigit()

            if is_digits:
                new_value = value
            else:
                new_value = str(value).upper()

            processed_data.append(tuple(row_list + [new_value]))

        base_name = save_file_name.split('.', 1)[0]
        new_file_name = base_name + '_process.xlsx'
        write_result = self.write_excel(processed_data, new_file_name)
        return (write_result, new_file_name)

import unittest
import os

class ExcelProcessorTestReadExcel(unittest.TestCase):
    def test_read_excel_1(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age', 'Country'],
                ['John', 25, 'USA'],
                ['Alice', 30, 'Canada'],
                ['Bob', 35, 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age', 'Country'),
            ('John', 25, 'USA'),
            ('Alice', 30, 'Canada'),
            ('Bob', 35, 'Australia')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_2(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Age'],
                ['John', 25],
                ['Alice', 30],
                ['Bob', 35]]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Age'),
            ('John', 25),
            ('Alice', 30),
            ('Bob', 35)
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_3(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name'],
                ['John'],
                ['Alice'],
                ['Bob']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name',),
            ('John',),
            ('Alice',),
            ('Bob',)
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_4(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Country'],
                ['John', 'USA'],
                ['Alice', 'Canada'],
                ['Bob', 'Australia']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Country'),
            ('John', 'USA'),
            ('Alice', 'Canada'),
            ('Bob', 'Australia')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_5(self):
        self.test_file_name = 'test_data.xlsx'
        data = [['Name', 'Country'],
                ['John', 'USA']]
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in data:
            sheet.append(row)
        workbook.save(self.test_file_name)
        workbook.close()

        processor = ExcelProcessor()
        data = processor.read_excel(self.test_file_name)
        expected_data = [
            ('Name', 'Country'),
            ('John', 'USA')
        ]
        self.assertEqual(data, expected_data)

    def test_read_excel_6(self):
        self.test_file_name = ''
        processor = ExcelProcessor()
        res = processor.read_excel(self.test_file_name)
        self.assertEqual(res, None)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
