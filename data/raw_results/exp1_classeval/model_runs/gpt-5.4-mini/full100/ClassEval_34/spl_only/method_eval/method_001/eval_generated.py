from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class DocFileHandler:
    def __init__(self, file_path):
        self.file_path = file_path

    def read_text(self):
        doc = Document(self.file_path)
        text = []
        for paragraph in doc.paragraphs:
            text.append(paragraph.text)
        return "\n".join(text)

    def write_text(self, content, font_size=12, alignment="left"):
        try:
            doc = Document()
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(content)
            font = run.font
            font.size = Pt(font_size)
            alignment_value = self._get_alignment_value(alignment)
            paragraph.alignment = alignment_value
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def add_heading(self, heading, level=1):
        try:
            doc = Document(self.file_path)
            doc.add_heading(heading, level=level)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def add_table(self, data):
        try:
            doc = Document(self.file_path)
            table = doc.add_table(rows=len(data), cols=len(data[0]))
            for i, row in enumerate(data):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = str(cell)
            doc.save(self.file_path)
            return True
        except Exception:
            return False

    def _get_alignment_value(self, alignment):
        alignment_options = {
            "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
            "center": WD_PARAGRAPH_ALIGNMENT.CENTER,
            "right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
        }
        resolved_alignment = alignment_options.get(alignment.lower(), WD_PARAGRAPH_ALIGNMENT.LEFT)
        return resolved_alignment

import unittest
import os

class DocFileHandlerTestWriteText(unittest.TestCase):
    def setUp(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

    def tearDown(self):
        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_write_text_1(self):
        new_content = "New content 1"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_2(self):
        new_content = "New content 2"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_3(self):
        new_content = "New content 3"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_4(self):
        new_content = "New content 4"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_5(self):
        new_content = ""
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
