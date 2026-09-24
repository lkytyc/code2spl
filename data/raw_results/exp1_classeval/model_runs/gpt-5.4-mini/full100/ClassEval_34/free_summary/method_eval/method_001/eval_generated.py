from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


class DocFileHandler:
    def __init__(self, file_path):
        self.file_path = file_path

    def _get_alignment_value(self, alignment):
        alignment_map = {
            'left': WD_PARAGRAPH_ALIGNMENT.LEFT,
            'center': WD_PARAGRAPH_ALIGNMENT.CENTER,
            'right': WD_PARAGRAPH_ALIGNMENT.RIGHT,
            'justify': WD_PARAGRAPH_ALIGNMENT.JUSTIFY,
        }
        return alignment_map.get(str(alignment).lower(), WD_PARAGRAPH_ALIGNMENT.LEFT)

    def read_text(self):
        document = Document(self.file_path)
        return '\n'.join(paragraph.text for paragraph in document.paragraphs)

    def write_text(self, content, font_size=12, alignment='left'):
        try:
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.alignment = self._get_alignment_value(alignment)
            run = paragraph.add_run(str(content))
            run.font.size = Pt(font_size)
            document.save(self.file_path)
            return True
        except Exception:
            return False

    def add_heading(self, heading, level=1):
        try:
            document = Document(self.file_path)
            document.add_heading(str(heading), level=level)
            document.save(self.file_path)
            return True
        except Exception:
            return False

    def add_table(self, data):
        try:
            document = Document(self.file_path)
            if not data:
                document.save(self.file_path)
                return True
            rows = len(data)
            cols = max(len(row) for row in data) if rows > 0 else 0
            table = document.add_table(rows=rows, cols=cols)
            for i, row in enumerate(data):
                for j in range(cols):
                    value = row[j] if j < len(row) else ""
                    table.cell(i, j).text = str(value)
            document.save(self.file_path)
            return True
        except Exception:
            return False

import unittest
import os

class DocFileHandlerTestWriteText(unittest.TestCase):
    def setUp(self):
        self.file_path = "test_example.docx"
        self.handler = DocFileHandler(self.file_path)
        doc = Document()
        doc.add_paragraph("Initial content")
        doc.save(self.file_path)

    def tearDown(self):
        if os.path.exists(self.file_path):
            os.remove(self.file_path)

    def test_write_text_1(self):
        new_content = "New content 1"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_2(self):
        new_content = "New content 2"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_3(self):
        new_content = "New content 3"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_4(self):
        new_content = "New content 4"
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

    def test_write_text_5(self):
        new_content = ""
        self.handler.write_text(new_content)
        text_content = self.handler.read_text()
        self.assertEqual(text_content, new_content)

if __name__ == '__main__':
    import unittest
    unittest.main(verbosity=2)
